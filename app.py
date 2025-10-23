# ====================================================================================
# app.py — Interview Master (수식/화학식 폴프루프 + 목록 안정화 + prompt.txt 캐시)
#  - prompt.txt: 최초 1회만 읽어 캐시, {text_from_pdf} 치환
#  - PDF: 인라인 $...$ → mathtext PNG, 화학식/수식 라인(⇌, →, ₊, ⁻, ₀-₉ 등) → Pillow 유니코드 PNG
#  - DOCX: 숫자/불릿 목록 안정 파서, 제목 직후 쓰레기 블릿/빈 문단 제거
#  - 수식/텍스트 이미지 렌더링 결과 캐시(lru_cache)
# ====================================================================================

import os, io, re, pdfplumber
from functools import lru_cache
from flask import Flask, request, jsonify, send_file, Response
from flask_cors import CORS
from dotenv import load_dotenv
import google.generativeai as genai

# ===== ReportLab (PDF)
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm

# ===== python-docx (DOCX)
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

# ===== matplotlib for LaTeX (inline math)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
matplotlib.rcParams['mathtext.fontset'] = 'dejavusans'

# ===== Pillow for Unicode chemical/equation lines =====
from PIL import Image as PILImage, ImageDraw, ImageFont

# ------------------------------------------------------------------------------------
# 앱/모델 설정
# ------------------------------------------------------------------------------------
load_dotenv()
app = Flask(__name__)
CORS(app)

try:
    genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))
    model = genai.GenerativeModel('gemini-2.5-flash')
except Exception as e:
    print(f"[WARN] Gemini 설정 실패: {e}")
    model = None

# ------------------------------------------------------------------------------------
# 유틸
# ------------------------------------------------------------------------------------
_ctrl_pattern = re.compile(r'[\x00-\x1F\x7F-\x9F]')

def sanitize_text(t: str) -> str:
    """제어문자/눈에 보이는 개행 화살표 제거"""
    return _ctrl_pattern.sub('', t).replace('↩', '').replace('\u2028', '\n')

def escape_html(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

# 인라인 수식/블록 수식 탐지
_inline_math_re = re.compile(r'\$(.+?)\$')
_block_math_single_re = re.compile(r'^\s*\$\$(.+?)\$\$\s*$')
_block_math_open_re   = re.compile(r'^\s*\$\$(.+)$')
_block_math_close_re  = re.compile(r'^(.+)\$\$\s*$')

# 화학식/수식 "의심" 라인(유니코드로 그리기): 화살표/첨자/수학기호/반응식 구조 등
_equation_like_re = re.compile(
    r'(⇌|→|←|↔|≈|≤|≥|±|√|∞|∑|∫|⊂|⊃|⇒|⇔|↦|↣|→|⇄|⇆|⇋|⇌|⇝|⇥|⇤|⇧|⇩|↕|⟶)'
    r'|[\u2070-\u209F]'         # super/subscripts
    r'|[₀-₉⁰-⁹⁺⁻]'              # more subs/supers/dingbats
    r'|[A-Za-z]\d'              # X2 형태
    r'|\[[^\]]+\]\s*[0-9]*[⁺⁻]?' # [CoCl4]2- 등
)

def is_equation_like_line(s: str) -> bool:
    return bool(_equation_like_re.search(s))

def is_trash_line(s: str) -> bool:
    """빈 줄/빈 블릿/쓰레기 기호 줄 제거"""
    st = s.strip()
    if st == "": return True
    if st in ('•', '-', '*', '+'): return True
    if re.fullmatch(r'[•\-\*\+]+', st): return True
    if re.fullmatch(r'(\-|\*|\+)\s*', st): return True
    return False

# ------------------------------------------------------------------------------------
# prompt.txt 읽기 (캐싱)
# ------------------------------------------------------------------------------------
@lru_cache(maxsize=1)
def load_prompt_template() -> str:
    """같은 폴더의 prompt.txt를 최초 1회만 읽어 캐시"""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    prompt_path = os.path.join(script_dir, 'prompt.txt')
    try:
        with open(prompt_path, 'r', encoding='utf-8') as f:
            tpl = f.read()
            if '{text_from_pdf}' not in tpl:
                tpl = tpl.rstrip() + "\n\n<학생부 내용>\n{text_from_pdf}\n"
            return tpl
    except Exception as e:
        print(f"[WARN] prompt.txt 읽기 실패: {e}")
        return ("다음 학생부 내용을 분석하여 맞춤형 면접 질문과 모범답변을 "
                "마크다운으로 생성하세요.\n\n<학생부 내용>\n{text_from_pdf}\n")

def build_prompt(extracted_text: str) -> str:
    tpl = load_prompt_template()
    return tpl.replace('{text_from_pdf}', extracted_text)

# ------------------------------------------------------------------------------------
# 인라인 마크다운 → ReportLab 문단 태그
# ------------------------------------------------------------------------------------
_inline_code_re   = re.compile(r'`([^`]+)`')
_inline_bold_re   = re.compile(r'\*\*(.+?)\*\*|__(.+?)__')
_inline_italic_re = re.compile(r'(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)|_(.+?)_')
_fallback_re      = re.compile(r'([\u0370-\u03FF\u1F00-\u1FFF\u2070-\u209F\u2200-\u22FF]+)')

def md_inline_to_rl(text: str, fallback_font: str = 'DejaVuSans') -> str:
    """ReportLab 문단용 마크다운 변환 + 폰트 폴백 래핑"""
    code_spans = []
    def _save(m):
        code_spans.append(m.group(1))
        return f"\u0000CODE{len(code_spans)-1}\u0000"
    text = _inline_code_re.sub(_save, text)
    t = escape_html(text)
    t = _fallback_re.sub(lambda m: f"<font name='{fallback_font}'>{m.group(1)}</font>", t)
    t = _inline_bold_re.sub(lambda m: f"<b>{m.group(1) or m.group(2) or ''}</b>", t)
    t = _inline_italic_re.sub(lambda m: f"<i>{m.group(1) or m.group(2) or ''}</i>", t)
    t = re.sub(r'\u0000CODE(\d+)\u0000',
               lambda m: f"<font name='Courier'>{escape_html(code_spans[int(m.group(1))])}</font>", t)
    return t

# ------------------------------------------------------------------------------------
# 라텍스/유니코드 렌더링 (캐시)
# ------------------------------------------------------------------------------------
@lru_cache(maxsize=256)
def render_latex_png(latex: str, dpi: int = 230, pad: float = 0.12) -> bytes:
    """라텍스 수식을 PNG로 렌더링"""
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.patch.set_alpha(0.0)
    ax = fig.add_subplot(111)
    ax.axis("off")
    ax.text(0.5, 0.5, f"${latex}$", fontsize=16, ha="center", va="center")
    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=dpi, transparent=True, bbox_inches="tight", pad_inches=pad)
    plt.close(fig)
    buf.seek(0)
    return buf.read()

@lru_cache(maxsize=256)
def render_unicode_text_png(text: str, font_size: int = 16, pad: int = 6) -> bytes:
    """유니코드 라인(화학식/수식)을 Pillow로 그려 PNG 생성 (폰트: DejaVuSans → NanumGothic 폴백)"""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    dejavu = os.path.join(script_dir, 'DejaVuSans.ttf')
    nanum  = os.path.join(script_dir, 'NanumGothic.ttf')
    font_path = dejavu if os.path.exists(dejavu) else nanum
    font = ImageFont.truetype(font_path, font_size)

    # 텍스트 bbox 계산
    dummy = PILImage.new("RGBA", (10, 10), (0, 0, 0, 0))
    d = ImageDraw.Draw(dummy)
    bbox = d.textbbox((0, 0), text, font=font)
    w, h = bbox[2] - bbox[0], bbox[3] - bbox[1]
    img = PILImage.new("RGBA", (w + 2*pad, h + 2*pad), (0, 0, 0, 0))
    d2 = ImageDraw.Draw(img)
    d2.text((pad, pad), text, font=font, fill=(0, 0, 0, 255))
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out.read()

# ------------------------------------------------------------------------------------
# 생성 API (스트리밍)
# ------------------------------------------------------------------------------------
@app.route("/generate", methods=["POST"])
def generate_interview_questions_stream():
    if model is None:
        return Response("AI 모델이 초기화되지 않았습니다.", status=500, mimetype="text/plain")
    if "file" not in request.files:
        return Response("요청에 파일이 없습니다.", status=400, mimetype="text/plain")
    file = request.files["file"]
    if file.filename == "":
        return Response("파일이 선택되지 않았습니다.", status=400, mimetype="text/plain")

    try:
        with pdfplumber.open(file.stream) as pdf:
            extracted_text = "\n".join(p.extract_text() or "" for p in pdf.pages)

        if len(extracted_text.strip()) < 50:
            return Response("PDF에서 충분한 텍스트를 추출하지 못했습니다.", status=400, mimetype="text/plain")

        prompt = build_prompt(extracted_text)

        def stream():
            try:
                for chunk in model.generate_content(prompt, stream=True):
                    yield chunk.text or ""
            except Exception as e:
                yield f"\n[스트리밍 오류] {e}"

        return Response(stream(), mimetype="text/plain; charset=utf-8")
    except Exception as e:
        return Response(f"파일 처리 오류: {e}", status=500, mimetype="text/plain")

# ------------------------------------------------------------------------------------
# 다운로드 API
# ------------------------------------------------------------------------------------
@app.route("/download", methods=["POST"])
def download_file():
    fmt = request.args.get("format") or request.form.get("format")
    content = request.form.get("content")
    if not content:
        return jsonify({"error": "전송된 내용이 없습니다. (content 누락)"}), 400

    content = sanitize_text(content)
    if fmt == "pdf":
        buf, mime, name = create_pdf_from_markdown(content)
    elif fmt in ("docx", "doc"):
        buf, mime, name = create_docx_from_markdown(content)
    else:
        return jsonify({"error": "지원하지 않는 형식입니다. (pdf, docx)"}), 400

    return send_file(buf, as_attachment=True, download_name=name, mimetype=mime)

# ------------------------------------------------------------------------------------
# PDF 변환 (목록/제목/코드 + 인라인/블록 수식 + 화학식 유니코드 이미지)
# ------------------------------------------------------------------------------------
def _append_text_with_inline_math_pdf(flow, text, style, fallback_font):
    """한 줄 내 $...$ 인라인 수식을 이미지로 끼워 넣어 Paragraph들과 섞어 추가"""
    pos = 0
    has_any = False
    for m in _inline_math_re.finditer(text):
        has_any = True
        left = text[pos:m.start()]
        if left.strip():
            flow.append(Paragraph(md_inline_to_rl(left, fallback_font), style))
        latex = m.group(1).strip()
        img_bytes = render_latex_png(latex)
        flow.append(Image(io.BytesIO(img_bytes), width=120, preserveAspectRatio=True))
        pos = m.end()
    tail = text[pos:]
    if has_any and tail.strip():
        flow.append(Paragraph(md_inline_to_rl(tail, fallback_font), style))
    return has_any

def create_pdf_from_markdown(md_text: str):
    buf = io.BytesIO()
    try:
        # 폰트 등록
        script_dir = os.path.dirname(os.path.abspath(__file__))
        nanum = os.path.join(script_dir, 'NanumGothic.ttf')
        dejavu = os.path.join(script_dir, 'DejaVuSans.ttf')
        pdfmetrics.registerFont(TTFont('NanumGothic', nanum))
        fallback_font = 'NanumGothic'
        if os.path.exists(dejavu):
            pdfmetrics.registerFont(TTFont('DejaVuSans', dejavu))
            fallback_font = 'DejaVuSans'

        # 스타일
        base = getSampleStyleSheet()
        body = ParagraphStyle('Body', parent=base['Normal'], fontName='NanumGothic',
                              fontSize=11, leading=16, spaceAfter=6)
        h1 = ParagraphStyle('H1', parent=body, fontSize=18, leading=24, spaceBefore=8, spaceAfter=10)
        h2 = ParagraphStyle('H2', parent=body, fontSize=15, leading=21, spaceBefore=6, spaceAfter=8)
        h3 = ParagraphStyle('H3', parent=body, fontSize=13, leading=19, spaceBefore=4, spaceAfter=6)
        code_style = ParagraphStyle('Code', parent=body, fontName='Courier',
                                    backColor='#f4f7fb', leftIndent=6, rightIndent=6,
                                    leading=15, spaceBefore=4, spaceAfter=8)

        flow = []
        lines = md_text.replace('\r\n', '\n').replace('\r', '\n').split('\n')

        # 목록 누적기
        def flush_list(lst, numbered=False):
            if not lst: return
            lf = ListFlowable(
                [ListItem(Paragraph(md_inline_to_rl(item, fallback_font), body)) for item in lst],
                bulletType='1' if numbered else 'bullet', start='1', leftIndent=12
            )
            flow.append(lf)
            lst.clear()

        ul, ol = [], []
        in_code = False
        code_lines = []
        in_block_math = False
        math_acc = []
        skip_after_header = 0  # 제목 직후 2~3줄 쓰레기 방지

        for raw in lines:
            line = raw.rstrip()

            # 코드 토글
            if line.strip().startswith('```'):
                if in_code:
                    code_html = "<br/>".join(escape_html(x) for x in code_lines) or " "
                    flow.append(Paragraph(code_html, code_style))
                    code_lines.clear()
                    in_code = False
                else:
                    flush_list(ul); flush_list(ol, numbered=True)
                    in_code = True
                continue
            if in_code:
                code_lines.append(line); continue

            # 쓰레기/빈 라인 스킵
            if is_trash_line(line):
                if skip_after_header > 0:  # 제목 직후 무의미 라인 무시
                    continue
                # 일반 빈 줄은 여백만
                flush_list(ul); flush_list(ol, numbered=True)
                flow.append(Spacer(1, 2 * mm))
                continue

            # 블록 수식 ($$ ... $$)
            if not in_block_math and _block_math_single_re.match(line):
                flush_list(ul); flush_list(ol, numbered=True)
                latex = _block_math_single_re.match(line).group(1).strip()
                flow.append(Image(io.BytesIO(render_latex_png(latex)), width=150, preserveAspectRatio=True))
                continue
            if not in_block_math and _block_math_open_re.match(line):
                flush_list(ul); flush_list(ol, numbered=True)
                in_block_math = True
                math_acc.append(_block_math_open_re.match(line).group(1))
                continue
            if in_block_math:
                if _block_math_close_re.match(line):
                    math_acc.append(_block_math_close_re.match(line).group(1))
                    latex = '\n'.join(math_acc).strip()
                    flow.append(Image(io.BytesIO(render_latex_png(latex)), width=150, preserveAspectRatio=True))
                    math_acc.clear(); in_block_math = False
                else:
                    math_acc.append(line)
                continue

            # 제목
            if line.startswith('### '):
                flush_list(ul); flush_list(ol, numbered=True)
                flow.append(Paragraph(md_inline_to_rl(line[4:].strip(), fallback_font), h3))
                skip_after_header = 3; continue
            if line.startswith('## '):
                flush_list(ul); flush_list(ol, numbered=True)
                flow.append(Paragraph(md_inline_to_rl(line[3:].strip(), fallback_font), h2))
                skip_after_header = 3; continue
            if line.startswith('# '):
                flush_list(ul); flush_list(ol, numbered=True)
                flow.append(Paragraph(md_inline_to_rl(line[2:].strip(), fallback_font), h1))
                skip_after_header = 3; continue

            # 목록
            if re.match(r'^(\-|\*|\+)\s+', line):
                text = line[2:].strip()
                if text:
                    ul.append(text)
                continue
            if re.match(r'^\d+\.\s+', line):
                text = re.sub(r'^\d+\.\s+', '', line).strip()
                if text:
                    ol.append(text)
                continue

            # 일반 문단
            flush_list(ul); flush_list(ol, numbered=True)

            # 인라인 수식이 있으면 섞어서 추가
            if _inline_math_re.search(line):
                handled = _append_text_with_inline_math_pdf(flow, line, body, fallback_font)
                if handled:
                    if skip_after_header > 0: skip_after_header -= 1
                    continue

            # 화학식/수식처럼 보이면 유니코드 텍스트 이미지를 사용
            if is_equation_like_line(line):
                img_bytes = render_unicode_text_png(line, font_size=15)
                flow.append(Image(io.BytesIO(img_bytes), width=170, preserveAspectRatio=True))
            else:
                flow.append(Paragraph(md_inline_to_rl(line, fallback_font), body))

            if skip_after_header > 0:
                skip_after_header -= 1

        # 마무리
        flush_list(ul); flush_list(ol, numbered=True)
        if in_code and code_lines:
            code_html = "<br/>".join(escape_html(x) for x in code_lines) or " "
            flow.append(Paragraph(code_html, code_style))

        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=20*mm, rightMargin=20*mm,
                                topMargin=17*mm, bottomMargin=17*mm)
        doc.build(flow)
        buf.seek(0)
        return buf, "application/pdf", "면접_질문+답변.pdf"
    except Exception as e:
        print(f"[ERR] PDF 생성 실패: {e}")
        raise

# ------------------------------------------------------------------------------------
# DOCX 변환 (목록/제목 안정 + 인라인/블록 수식 + 화학식 유니코드 이미지)
# ------------------------------------------------------------------------------------
_inline_token = re.compile(r'(\*\*.+?\*\*|__.+?__|`[^`]+`|\*(?!\*)(.+?)(?<!\*)\*|_.+?_)')

def _apply_run_font(run):
    run.font.size = Pt(11)
    run.font.name = '맑은 고딕'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

def _add_inline_runs(p, text):
    pos = 0
    for m in _inline_token.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()]); _apply_run_font(r)
        token = m.group(0)
        if token.startswith('**') or token.startswith('__'):
            r = p.add_run(token[2:-2]); r.bold = True; _apply_run_font(r)
        elif token.startswith('`') and token.endswith('`'):
            r = p.add_run(token[1:-1]); r.font.name = 'Consolas'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕'); r.font.size = Pt(11)
        else:
            r = p.add_run(token[1:-1]); r.italic = True; _apply_run_font(r)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:]); _apply_run_font(r)

def _add_math_image_run(p, latex):
    width_inch = min(0.8 + 0.05 * len(latex), 3.5)
    img = render_latex_png(latex)
    p.add_run().add_picture(io.BytesIO(img), width=Inches(width_inch))

def _add_unicode_line_image(p, text):
    img = render_unicode_text_png(text, font_size=15)
    p.add_run().add_picture(io.BytesIO(img), width=Inches(2.6))

def _tune_paragraph(p):
    pf = p.paragraph_format
    pf.line_spacing = 1.3
    pf.space_before = Pt(2)
    pf.space_after = Pt(3)

def create_docx_from_markdown(md_text: str):
    buf = io.BytesIO()
    try:
        doc = Document()
        style = doc.styles['Normal']
        style.font.name = '맑은 고딕'
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

        lines = md_text.replace('\r\n', '\n').replace('\r', '\n').split('\n')

        ul, ol = [], []
        def flush_lists():
            nonlocal ul, ol
            for it in ul:
                p = doc.add_paragraph(style='List Bullet'); _tune_paragraph(p)
                _add_inline_runs(p, it)
            for it in ol:
                p = doc.add_paragraph(style='List Number'); _tune_paragraph(p)
                _add_inline_runs(p, it)
            ul, ol = [], []

        in_code = False
        code_lines = []
        in_block_math = False
        math_acc = []
        skip_after_header = 0

        for raw in lines:
            line = raw.rstrip()

            # 코드
            if line.strip().startswith('```'):
                if in_code:
                    p = doc.add_paragraph(); _tune_paragraph(p)
                    for cl in code_lines:
                        r = p.add_run(cl + '\n')
                        r.font.name = 'Consolas'
                        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
                        r.font.size = Pt(10)
                    code_lines.clear(); in_code = False
                else:
                    flush_lists()
                    in_code = True
                continue
            if in_code:
                code_lines.append(line); continue

            # 빈/쓰레기 라인
            if is_trash_line(line):
                if skip_after_header > 0:
                    continue
                flush_lists()
                doc.add_paragraph()  # 적당한 여백
                continue

            # 블록 수식
            if not in_block_math and _block_math_single_re.match(line):
                flush_lists()
                p = doc.add_paragraph(); _tune_paragraph(p)
                _add_math_image_run(p, _block_math_single_re.match(line).group(1).strip())
                continue
            if not in_block_math and _block_math_open_re.match(line):
                flush_lists()
                in_block_math = True
                math_acc.append(_block_math_open_re.match(line).group(1))
                continue
            if in_block_math:
                if _block_math_close_re.match(line):
                    math_acc.append(_block_math_close_re.match(line).group(1))
                    p = doc.add_paragraph(); _tune_paragraph(p)
                    _add_math_image_run(p, '\n'.join(math_acc).strip())
                    math_acc.clear(); in_block_math = False
                else:
                    math_acc.append(line)
                continue

            # 제목
            if line.startswith('### '):
                flush_lists()
                p = doc.add_paragraph(line[4:].strip(), style='Heading 3'); _tune_paragraph(p)
                skip_after_header = 3; continue
            if line.startswith('## '):
                flush_lists()
                p = doc.add_paragraph(line[3:].strip(), style='Heading 2'); _tune_paragraph(p)
                skip_after_header = 3; continue
            if line.startswith('# '):
                flush_lists()
                p = doc.add_paragraph(line[2:].strip(), style='Heading 1'); _tune_paragraph(p)
                skip_after_header = 3; continue

            # 목록
            if re.match(r'^(\-|\*|\+)\s+', line):
                text = line[2:].strip()
                if text:
                    ul.append(text)
                continue
            if re.match(r'^\d+\.\s+', line):
                text = re.sub(r'^\d+\.\s+', '', line).strip()
                if text:
                    ol.append(text)
                continue

            # 일반 문단
            flush_lists()
            p = doc.add_paragraph(); _tune_paragraph(p)

            # 인라인 수식
            cursor = 0
            had_inline = False
            for m in _inline_math_re.finditer(line):
                had_inline = True
                pre = line[cursor:m.start()]
                if pre: _add_inline_runs(p, pre)
                _add_math_image_run(p, m.group(1).strip())
                cursor = m.end()
            tail = line[cursor:]
            if tail:
                # 화학식/수식처럼 보이는 일반 텍스트는 유니코드 이미지로
                if is_equation_like_line(tail) and not had_inline:
                    _add_unicode_line_image(p, tail)
                else:
                    _add_inline_runs(p, tail)

            if skip_after_header > 0:
                skip_after_header -= 1

        # 마무리
        flush_lists()
        if in_code and code_lines:
            p = doc.add_paragraph(); _tune_paragraph(p)
            for cl in code_lines:
                r = p.add_run(cl + '\n')
                r.font.name = 'Consolas'
                r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
                r.font.size = Pt(10)

        if len(doc.paragraphs) == 0:
            t = doc.add_paragraph('AI 생성 면접 질문 및 모범 답변', style='Heading 1'); _tune_paragraph(t)

        doc.save(buf)
        buf.seek(0)
        return buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "면접_질문+답변.docx"
    except Exception as e:
        print(f"[ERR] DOCX 생성 실패: {e}")
        raise

# ------------------------------------------------------------------------------------
@app.get("/healthz")
def healthz(): return "ok", 200

@app.get("/")
def root(): return "AI Interview Master backend running", 200
